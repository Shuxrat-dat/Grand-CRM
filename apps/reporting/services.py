import io
from decimal import Decimal
from datetime import date, datetime
from django.db.models import Sum, Count, Q
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from apps.billing.models import MonthlyCharge, Payment, PaymentAllocation
from apps.payroll.models import SalaryRecord, SalaryPayment
from apps.finance.models import Expense
from apps.students.models import Student
from apps.academics.models import StudyGroup, Course, Teacher
from apps.attendance.models import Lesson, AttendanceRecord


def get_dashboard_metrics(start_date: date, end_date: date, is_owner: bool):
    """
    Computes key metrics for dashboard based on user role and date filter.
    Everything calculated from real database records.
    """
    # Active students count
    active_students_count = Student.objects.filter(status=Student.Status.ACTIVE).count()
    active_groups_count = StudyGroup.objects.filter(is_active=True).count()
    active_teachers_count = Teacher.objects.filter(status=Teacher.Status.ACTIVE).count()

    # Cash Inflows in period
    payments_in_period = Payment.objects.filter(
        status=Payment.Status.COMPLETED,
        paid_at__date__gte=start_date,
        paid_at__date__lte=end_date
    ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    # Accrued charges in period
    charges_in_period = MonthlyCharge.objects.exclude(
        status=MonthlyCharge.Status.CANCELLED
    ).filter(
        billing_month__gte=start_date,
        billing_month__lte=end_date
    ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    # Total outstanding student debt across the system
    all_active_charges = MonthlyCharge.objects.exclude(
        status__in=[MonthlyCharge.Status.PAID, MonthlyCharge.Status.CANCELLED]
    )
    total_debt = sum([c.remaining_balance for c in all_active_charges], Decimal('0.00'))

    # Overdue charges count
    today = timezone.now().date()
    overdue_count = MonthlyCharge.objects.filter(
        status=MonthlyCharge.Status.OVERDUE
    ).count()

    metrics = {
        'active_students_count': active_students_count,
        'active_groups_count': active_groups_count,
        'active_teachers_count': active_teachers_count,
        'payments_in_period': payments_in_period,
        'charges_in_period': charges_in_period,
        'total_debt': total_debt,
        'overdue_count': overdue_count,
        'start_date': start_date,
        'end_date': end_date,
    }

    if is_owner:
        # Accrued Salaries in period
        salaries_accrued = SalaryRecord.objects.exclude(
            status=SalaryRecord.Status.CANCELLED
        ).filter(
            period_month__gte=start_date,
            period_month__lte=end_date
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        # Actual salary payments disbursed in period
        salaries_paid = SalaryPayment.objects.filter(
            status=SalaryPayment.Status.COMPLETED,
            paid_at__date__gte=start_date,
            paid_at__date__lte=end_date
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        # Operating expenses in period
        expenses_total = Expense.objects.filter(
            status=Expense.Status.ACTIVE,
            occurred_at__gte=start_date,
            occurred_at__lte=end_date
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        # Outstanding salary debt
        all_active_salaries = SalaryRecord.objects.exclude(
            status__in=[SalaryRecord.Status.PAID, SalaryRecord.Status.CANCELLED]
        )
        total_salary_debt = sum([s.remaining_balance for s in all_active_salaries], Decimal('0.00'))

        # Accrual profit = Accrued Revenue - Accrued Salaries - Operating Expenses
        accrual_profit = charges_in_period - salaries_accrued - expenses_total

        # Net cash flow = Cash Inflows - Cash Salary Payments - Cash Expenses
        net_cash_flow = payments_in_period - salaries_paid - expenses_total

        metrics.update({
            'salaries_accrued': salaries_accrued,
            'salaries_paid': salaries_paid,
            'expenses_total': expenses_total,
            'total_salary_debt': total_salary_debt,
            'accrual_profit': accrual_profit,
            'net_cash_flow': net_cash_flow,
        })

    return metrics


def get_cash_flow_report(start_date: date, end_date: date):
    """
    Cash Flow Statement (Kassa / Pul oqimi hisoboti).
    """
    payments = Payment.objects.filter(
        status=Payment.Status.COMPLETED,
        paid_at__date__gte=start_date,
        paid_at__date__lte=end_date
    )
    inflows_total = payments.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    inflows_by_method = payments.values('method').annotate(total=Sum('amount')).order_by('-total')

    salary_payments = SalaryPayment.objects.filter(
        status=SalaryPayment.Status.COMPLETED,
        paid_at__date__gte=start_date,
        paid_at__date__lte=end_date
    )
    salaries_total = salary_payments.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    expenses = Expense.objects.filter(
        status=Expense.Status.ACTIVE,
        occurred_at__gte=start_date,
        occurred_at__lte=end_date
    )
    expenses_total = expenses.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    expenses_by_category = expenses.values('category').annotate(total=Sum('amount')).order_by('-total')

    outflows_total = salaries_total + expenses_total
    net_cash_flow = inflows_total - outflows_total

    return {
        'start_date': start_date,
        'end_date': end_date,
        'inflows_total': inflows_total,
        'inflows_by_method': inflows_by_method,
        'salaries_total': salaries_total,
        'expenses_total': expenses_total,
        'expenses_by_category': expenses_by_category,
        'outflows_total': outflows_total,
        'net_cash_flow': net_cash_flow,
    }


def get_accrual_profit_report(start_date: date, end_date: date):
    """
    Accrual Basis P&L Statement (Hisoblangan operatsion foyda va zarar).
    """
    charges = MonthlyCharge.objects.exclude(
        status=MonthlyCharge.Status.CANCELLED
    ).filter(
        billing_month__gte=start_date,
        billing_month__lte=end_date
    )
    revenue_accrued = charges.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    salaries = SalaryRecord.objects.exclude(
        status=SalaryRecord.Status.CANCELLED
    ).filter(
        period_month__gte=start_date,
        period_month__lte=end_date
    )
    salaries_accrued = salaries.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    expenses = Expense.objects.filter(
        status=Expense.Status.ACTIVE,
        occurred_at__gte=start_date,
        occurred_at__lte=end_date
    )
    expenses_accrued = expenses.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    total_costs = salaries_accrued + expenses_accrued
    calculated_profit = revenue_accrued - total_costs

    return {
        'start_date': start_date,
        'end_date': end_date,
        'revenue_accrued': revenue_accrued,
        'salaries_accrued': salaries_accrued,
        'expenses_accrued': expenses_accrued,
        'total_costs': total_costs,
        'calculated_profit': calculated_profit,
    }


def get_debtors_report():
    """
    List of students with outstanding debt balance.
    """
    students = Student.objects.filter(
        status__in=[Student.Status.ACTIVE, Student.Status.SUSPENDED]
    ).prefetch_related('enrollments__group', 'monthly_charges')

    debtors = []
    for s in students:
        debt = s.total_debt
        if debt > Decimal('0.00'):
            overdue_count = s.monthly_charges.filter(status=MonthlyCharge.Status.OVERDUE).count()
            active_groups = ", ".join([e.group.name for e in s.active_enrollments])
            debtors.append({
                'student': s,
                'active_groups': active_groups or "Guruhsiz",
                'total_charged': s.total_charged,
                'total_paid': s.total_allocated_payments,
                'total_debt': debt,
                'overdue_count': overdue_count,
            })

    debtors.sort(key=lambda x: x['total_debt'], reverse=True)
    return debtors


def export_report_to_excel(title: str, headers: list, rows: list, total_row: list = None) -> io.BytesIO:
    """
    Generates an Excel workbook using openpyxl.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hisobot"

    # Title styling
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Headers styling
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="334155")
    thin_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )

    ws.append([]) # empty row 2
    ws.append(headers) # row 3

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[3].height = 24

    # Data rows
    row_idx = 4
    for r in rows:
        ws.append(r)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            if isinstance(cell.value, (int, float, Decimal)):
                cell.alignment = Alignment(horizontal="right")
        row_idx += 1

    # Total row if provided
    if total_row:
        ws.append(total_row)
        total_font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
        total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = thin_border
            if isinstance(cell.value, (int, float, Decimal)):
                cell.alignment = Alignment(horizontal="right")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_report_to_pdf(title: str, headers: list, rows: list, summary_text: str = "") -> io.BytesIO:
    """
    Generates a PDF document using ReportLab.
    """
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()

    elements = []
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1, # Center
        spaceAfter=15
    )
    elements.append(Paragraph(title, title_style))

    if summary_text:
        summary_style = ParagraphStyle(
            'DocSummary',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#475569'),
            spaceAfter=15
        )
        elements.append(Paragraph(summary_text, summary_style))

    # Format table data
    table_data = [headers]
    for row in rows:
        formatted_row = [str(item) for item in row]
        table_data.append(formatted_row)

    col_widths = None # auto
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
    ]))

    elements.append(table)
    doc.build(elements)
    output.seek(0)
    return output
